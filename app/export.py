"""Builds the Excel (.xlsx) export."""
from datetime import datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import catalogue
from .config import settings

GREEN = "1C5841"
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=GREEN)
STATUS_LABELS = {"new": "New", "reviewed": "Reviewed", "actioned": "Actioned", "archived": "Archived"}
RETURN_LABELS = {"yes": "Yes", "maybe": "Maybe", "no": "No"}
LANGUAGE_LABELS = {"en": "English", "ur": "Urdu", "roman": "Roman Urdu"}


def _header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def build_workbook(rows: list[dict], stats: dict, filters_label: str) -> bytes:
    tz = ZoneInfo(settings.timezone)
    wb = Workbook()

    # -------------------------------------------------------- Responses
    ws = wb.active
    ws.title = "Responses"
    questions = [q for q in catalogue.QUESTIONS if q["kind"] in catalogue.ANSWER_KINDS]
    headers = ["Reference", f"Submitted ({settings.timezone})", "Status", "Branch", "Order type",
               "Food (1-5)", "Would come back", "Needs attention", "Issues"]
    headers += [catalogue.title(q) + (" (1-5)" if q["kind"] == "stars" else "") for q in questions]
    headers += ["Other item ordered", "Comments", "Wants offers", "Name", "Mobile", "Language", "Admin note"]
    ws.append(headers)
    first_free_text = len(headers) - 6  # "Other item ordered" onward is typed by people

    for r in rows:
        answers = r["answers"]
        line = [
            r["reference_code"], r["submitted_at"].astimezone(tz).replace(tzinfo=None),
            STATUS_LABELS.get(r["status"], r["status"]), r["branch_name"], r["service_name"],
            r["food_rating"], RETURN_LABELS.get(r["would_return"]), "Yes" if r["needs_attention"] else "No",
            "; ".join(r["issues"]),
        ]
        for q in questions:
            v = answers.get(q["key"])
            line.append(v if q["kind"] == "stars" else (catalogue.label(q, v) or None))
        line += [r["other_item"], r["comments"], "Yes" if r["wants_offers"] else "No",
                 r["customer_name"], r["customer_phone"], LANGUAGE_LABELS.get(r["language"], r["language"]),
                 r["admin_note"]]
        ws.append(line)
        # openpyxl treats any string starting with "=" as a formula; keep typed text as text.
        for col in range(first_free_text, len(headers) + 1):
            cell = ws.cell(row=ws.max_row, column=col)
            if isinstance(cell.value, str):
                cell.data_type = "s"

    _header(ws)
    ws.row_dimensions[1].height = 60
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(ws.max_row, 1)}"
    widths = {1: 20, 2: 18, 3: 11, 4: 26, 5: 13, 6: 10, 7: 12, 8: 12, 9: 34}
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(i, 22)
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "yyyy-mm-dd hh:mm"

    # ---------------------------------------------------------- Summary
    sm = wb.create_sheet("Summary")
    sm.append(["Sultan Yakhni Pulao customer feedback"])
    sm["A1"].font = Font(bold=True, size=14, color=GREEN)
    sm.append(["Exported", datetime.now(tz).replace(tzinfo=None)])
    sm["B2"].number_format = "yyyy-mm-dd hh:mm"
    sm.append(["Filters", filters_label])
    sm["B3"].data_type = "s"
    sm.append([])
    sm.append(["Responses", stats["total"]])
    sm.append(["Average food rating (1-5)", stats["avg_food"]])
    sm.append(["Would come back (%)", stats["return_pct"]])
    sm.append(["Needs attention", stats["attention"]])
    sm.append(["Signed up for offers", stats["offers"]])
    for r in range(5, 10):
        sm.cell(row=r, column=1).font = Font(bold=True)
    sm.append([])
    sm.append(["Branch", "Responses", "Food (1-5)", "Would come back (%)", "Needs attention"])
    _header(sm, sm.max_row)
    for b in stats["by_branch"]:
        sm.append([b["name"], b["total"], b["avg_food"], b["return_pct"], b["attention"]])
    sm.column_dimensions["A"].width = 34
    for col in "BCDE":
        sm.column_dimensions[col].width = 18

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
