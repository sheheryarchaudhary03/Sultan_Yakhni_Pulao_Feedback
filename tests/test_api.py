import re
from io import BytesIO

from openpyxl import load_workbook

DINE_IN = {
    "branch": "daska", "service": "dine_in", "food": 5, "staff_welcome": 4,
    "food_time": "10_20", "cleanliness": 5, "ambiance": "good", "dine_again": "definitely",
}
DELIVERY_BAD = {
    "branch": "johar-town", "service": "delivery", "food": 2, "on_time": "very_late",
    "arrived_hot": "cold", "delivery_complete": "missing_items", "rider_behaviour": "poor",
    "order_again": "no", "platform": "whatsapp", "ordering_system": "easy",
    "items": ["chicken_pulao", "shami", "other"], "favorite": "shami", "pulao_rating": "average",
}


def submit(client, answers, **extra):
    return client.post("/api/feedback", json={"language": "en", "answers": answers, "name": "Test customer", **extra})


def test_form_page_carries_catalogue(client):
    r = client.get("/")
    assert r.status_code == 200
    assert "{{CATALOGUE}}" not in r.text
    assert '"key":"food"' in r.text
    assert "script-src 'self'" in r.headers["content-security-policy"]


def test_submit_dine_in(client):
    r = submit(client, DINE_IN, comments="Great pulao")
    assert r.status_code == 201, r.text
    assert re.fullmatch(r"SYP-\d{6}-[A-Z0-9]{5}", r.json()["reference"])


def test_languages(client):
    for code in ("en", "ur", "roman"):
        r = client.post("/api/feedback", json={"language": code, "answers": DINE_IN, "name": "Test customer"})
        assert r.status_code == 201, (code, r.text)
    r = client.post("/api/feedback", json={"language": "fr", "answers": DINE_IN, "name": "Test customer"})
    assert r.status_code == 422


def test_missing_required_answer_is_refused(client):
    answers = {k: v for k, v in DINE_IN.items() if k != "cleanliness"}
    r = submit(client, answers)
    assert r.status_code == 422
    assert "cleanliness" in r.json()["detail"]


def test_unknown_option_and_question_are_refused(client):
    assert submit(client, {**DINE_IN, "ambiance": "fabulous"}).status_code == 422
    assert submit(client, {**DINE_IN, "colour": "red"}).status_code == 422
    assert submit(client, {**DINE_IN, "food": True}).status_code == 422
    assert submit(client, {**DINE_IN, "branch": "karachi"}).status_code == 422


def test_answers_for_hidden_questions_are_dropped(client, admin):
    # A delivery-only answer on a dine-in response is ignored, not stored.
    r = submit(client, {**DINE_IN, "rider_behaviour": "poor"})
    assert r.status_code == 201
    ref = r.json()["reference"]
    item = admin.get("/api/admin/responses", params={"q": ref}).json()["items"][0]
    detail = admin.get(f"/api/admin/responses/{item['id']}").json()
    keys = {row["key"] for s in detail["sections"] for row in s["rows"]}
    assert "rider_behaviour" not in keys
    assert item["needs_attention"] is False


def test_name_required_and_phone_checked(client, admin):
    r = submit(client, DINE_IN, name="")
    assert r.status_code == 422 and "name" in r.json()["detail"]
    assert submit(client, DINE_IN, phone="12345").status_code == 422
    r = submit(client, DINE_IN, name="Ayesha", phone="0321-7654321")
    assert r.status_code == 201
    item = admin.get("/api/admin/responses", params={"q": r.json()["reference"]}).json()["items"][0]
    assert item["customer_name"] == "Ayesha" and item["customer_phone"] == "0321-7654321"
    assert item["wants_offers"] is False


def test_offers_need_a_phone(client):
    r = submit(client, {**DINE_IN, "offers": "yes"})
    assert r.status_code == 422
    r = submit(client, {**DINE_IN, "offers": "yes"}, name="Ali", phone="0300 1234567")
    assert r.status_code == 201


def test_bad_delivery_needs_attention(client, admin):
    r = submit(client, DELIVERY_BAD, other_item="Kheer")
    assert r.status_code == 201, r.text
    items = admin.get("/api/admin/responses", params={"attention": "true", "branch": "johar-town"}).json()["items"]
    assert len(items) == 1
    issues = items[0]["issues"]
    for expected in ("Food rated 2/5", "Delivery very late", "Food arrived cold", "Items missing from order",
                     "Rider behaviour poor", "Would not come back"):
        assert expected in issues
    detail = admin.get(f"/api/admin/responses/{items[0]['id']}").json()
    assert detail["other_item"] == "Kheer"
    assert detail["branch_name"] == "Sultan Johar Town"
    rows = {row["key"]: row["value"] for s in detail["sections"] for row in s["rows"]}
    assert rows["items"] == "Chicken Yakhni Pulao, Shami Kabab, Other"
    assert rows["favorite"] == "Shami Kabab"


def test_honeypot_stores_nothing(client, admin):
    before = admin.get("/api/admin/stats").json()["total"]
    r = submit(client, DINE_IN, website="http://spam.example")
    assert r.status_code == 201
    assert admin.get("/api/admin/stats").json()["total"] == before


def test_admin_needs_sign_in(client):
    client.cookies.clear()
    assert client.get("/api/admin/stats").status_code == 401
    assert client.post("/api/admin/login", json={"username": "admin", "password": "nope"}).status_code == 401


def test_stats_and_breakdown(admin):
    s = admin.get("/api/admin/stats").json()
    assert s["total"] >= 4
    assert 1 <= s["avg_food"] <= 5
    assert {b["branch"] for b in s["by_branch"]} >= {"daska", "johar-town", "kashmir-park"}
    assert sum(x["n"] for x in s["food_dist"]) == s["total"]
    b = admin.get("/api/admin/breakdown/items").json()
    counts = {i["value"]: i["n"] for i in b["items"]}
    assert counts["shami"] == 1 and counts["other"] == 1
    assert admin.get("/api/admin/breakdown/food").json()["respondents"] == s["total"]
    assert admin.get("/api/admin/breakdown/nope").status_code == 404


def test_status_note_and_export(admin):
    item = admin.get("/api/admin/responses", params={"attention": "true"}).json()["items"][0]
    r = admin.patch(f"/api/admin/responses/{item['id']}", json={"status": "actioned", "admin_note": "Called, refunded"})
    assert r.status_code == 200
    assert r.json()["status"] == "actioned" and r.json()["admin_note"] == "Called, refunded"
    x = admin.get("/api/admin/export.xlsx")
    assert x.status_code == 200
    wb = load_workbook(BytesIO(x.content))
    assert wb.sheetnames == ["Responses", "Summary"]
    header = [c.value for c in wb["Responses"][1]]
    assert "Branch" in header and "How was the food? (1-5)" in header
    assert wb["Responses"].max_row >= 5


def test_change_password(admin):
    assert admin.get("/api/admin/me").json()["default_password"] is True
    r = admin.post("/api/admin/password", json={"current_password": "1234", "new_password": "pulao2026"})
    assert r.status_code == 200
    assert admin.get("/api/admin/me").json()["default_password"] is False
    admin.post("/api/admin/password", json={"current_password": "pulao2026", "new_password": "back2default"})
